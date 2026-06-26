#!/usr/bin/env python3
"""Generate assets/data/universities.json from KZ primary sources.

Sources (see sources.md):
  * egov.kz HTML   — canonical university registry: name, city, website (email).
  * Wikipedia      — university TYPE (national/state/private/...).
  * НЦТ XLSX        — minimum grant-competition scores per university x ГОП (2024).
  * НЦТ PDF        — ГОП -> two profile ЕНТ subjects (2025-26).

Universities come from egov (full list incl. non-grant private ones). Grant data
from the XLSX is JOINED onto egov universities by normalised name; XLSX unis that
don't match keep an `ovpo-*` fallback id so no grant data is lost.

Honesty: XLSX scores are `competition_min` (минимум для допуска к конкурсу), never
проходной балл. Every grant_threshold carries source_url. Type is set only on a
confident name match; otherwise left null.

Usage: python3 scrape.py [--egov f.html --wiki f.html --xlsx f.xlsx --pdf f.pdf]
Deps: requests, openpyxl, pypdf, beautifulsoup4
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import os
import re
import sys
import tempfile

import openpyxl
import pypdf
import requests
from bs4 import BeautifulSoup

EGOV_URL = "https://egov.kz/cms/ru/articles/2Fvusi_rk"
WIKI_URL = (
    "https://ru.wikipedia.org/wiki/"
    "Список_высших_учебных_заведений_Казахстана"
)
XLSX_URL = (
    "https://testcenter.kz/wp-content/uploads/2025/01/"
    "1-Минимальные-баллы-ЕНТ-на-Конкурс-2024-16.07.2024-1.xlsx"
)
PDF_URL = (
    "https://testcenter.kz/wp-content/uploads/2026/05/"
    "Список-специальностей-полной-формы-обучения-с-профильными-предметами.pdf"
)
GRANT_YEAR = 2024

SUBJECTS = sorted(
    [
        "Всемирная история", "Казахский язык", "Казахская литература",
        "Русский язык", "Русская литература", "Иностранный язык",
        "Основы права", "Творческий экзамен", "Биология", "География",
        "Математика", "Физика", "Информатика", "Химия",
    ],
    key=len, reverse=True,
)

FIELD_MAP = [
    ("педагог", "natural"), ("информац", "informatics"), ("информати", "informatics"),
    ("инженер", "engineering"), ("технич", "engineering"), ("производ", "engineering"),
    ("строит", "engineering"), ("транспорт", "engineering"),
    ("здравоохран", "medicine"), ("медиц", "medicine"), ("ветерин", "medicine"),
    ("естествен", "natural"), ("биолог", "natural"), ("сельск", "natural"),
    ("эконом", "economics"), ("бизнес", "economics"), ("управлен", "economics"),
    ("прав", "law"), ("юрид", "law"),
    ("искусств", "arts"), ("гуманитар", "arts"), ("социальн", "arts"),
    ("математ", "mathematics"), ("статист", "mathematics"),
]

# legal-form / filler words stripped before name matching
_STOP = [
    "некоммерческое акционерное общество", "акционерное общество",
    "товарищество с ограниченной ответственностью", "частное учреждение",
    "республиканское государственное предприятие на праве хозяйственного ведения",
    "республиканское государственное предприятие", "учреждение образования",
    "учреждение", "филиал", "имени", "им.", "высшего", "образования",
    "нао", "ао", "тоо", "рггп", "пхв",
]
_WIKI_TYPES = [
    ("национальн", "national"), ("международн", "international"),
    ("государственн", "state"), ("силов", "state"),
    ("акционирован", "autonomous"), ("частн", "private"),
    ("филиалы росс", "international"),
]


def _field_for(area_name):
    if not area_name:
        return None
    low = area_name.lower()
    for needle, key in FIELD_MAP:
        if needle in low:
            return key
    return None


def _norm_code(code):
    return code.strip().replace("В", "B").replace("М", "M").upper()


def _normalize(name):
    """Normalise a university name to a token set for matching."""
    s = " " + name.lower().replace("ё", "е") + " "
    s = s.replace("«", " ").replace("»", " ").replace('"', " ").replace("'", " ")
    s = re.sub(r"[^a-zа-я0-9 ]", " ", s)
    for w in _STOP:
        s = s.replace(" " + w + " ", " ")
    return frozenset(t for t in s.split() if len(t) > 1)


def _jaccard(a, b):
    if not a or not b:
        return 0.0
    return len(a & b) / len(a | b)


def _best_match(tokens, index, threshold=0.6):
    """Return the id of the best-matching entry in index {id: tokenset}."""
    best_id, best = None, 0.0
    for uid, toks in index.items():
        j = _jaccard(tokens, toks)
        if j > best:
            best, best_id = j, uid
    return best_id if best >= threshold else None


def _city_from_addr(addr):
    addr = (addr or "").replace("\xa0", " ")
    m = re.search(r"г\.?\s*([А-ЯЁ][А-Яа-яёЁ\-]+)", addr)
    return m.group(1) if m else None


def _website_from_email(email):
    email = (email or "").strip().lower().replace(" ", "")
    if not email:
        return None
    m = re.search(r"@([a-z0-9.\-]+)", email)
    dom = m.group(1) if m else None
    if not dom:
        m2 = re.search(r"([a-z0-9\-]+\.[a-z0-9.\-]+)",
                       email.replace("https://", "").replace("http://", ""))
        dom = m2.group(1) if m2 else None
    if not dom:
        return None
    return dom.strip("/").lstrip(".").replace("www.", "") or None


def _download(url, suffix):
    fd, path = tempfile.mkstemp(suffix=suffix)
    os.close(fd)
    resp = requests.get(url, timeout=90, headers={"User-Agent": "Mozilla/5.0"})
    resp.raise_for_status()
    with open(path, "wb") as fh:
        fh.write(resp.content)
    return path


# ── egov.kz: canonical universities ──────────────────────────────────────────
def parse_egov(path):
    soup = BeautifulSoup(open(path, encoding="utf-8").read(), "html.parser")
    unis, seen = [], {}
    for table in soup.find_all("table"):
        if table.find("table"):  # skip outer layout wrappers
            continue
        rows = table.find_all("tr")
        if not rows:
            continue
        head = rows[0].get_text(" ", strip=True)
        if "Наименование" not in head or "ОВПО" not in head:
            continue
        region_node = table.find_previous(string=re.compile("Высшие учебные заведения"))
        region = ""
        if region_node:
            region = re.sub(r".*заведения\s*", "", " ".join(region_node.split()))
            region = region.replace("г.", "").replace("области", "").strip()
        for r in rows[1:]:
            cells = r.find_all(["td", "th"])
            if len(cells) < 3:
                continue
            num = cells[0].get_text(" ", strip=True)
            name = cells[1].get_text(" ", strip=True)
            if not name or not re.match(r"^\d+$", num):
                continue
            addr = cells[2].get_text(" ", strip=True)
            email = cells[4].get_text(" ", strip=True) if len(cells) > 4 else ""
            website = _website_from_email(email)
            city = _city_from_addr(addr) or region or None
            base = website.split(".")[0] if website else None
            if not base or not re.match(r"^[a-z0-9\-]+$", base):
                base = f"egov-{len(unis) + 1}"
            uid = base
            n = 2
            while uid in seen:
                uid = f"{base}-{n}"
                n += 1
            seen[uid] = True
            unis.append({
                "id": uid, "name_ru": name, "city": city,
                "website": website, "source_url": EGOV_URL,
            })
    return unis


# ── Wikipedia: name -> type ──────────────────────────────────────────────────
# Each type section (h2 in a div.mw-heading) is followed by a
# <table class="sortable"> whose first column is the university name.
def parse_wiki_types(path):
    soup = BeautifulSoup(open(path, encoding="utf-8").read(), "html.parser")
    out = {}
    for heading in soup.find_all("div", class_=re.compile("mw-heading")):
        txt = heading.get_text(" ", strip=True).lower()
        current = None
        for needle, val in _WIKI_TYPES:
            if needle in txt:
                current = val
                break
        if not current:
            continue
        for sib in heading.find_next_siblings():
            classes = sib.get("class") or []
            if sib.name == "div" and any("mw-heading" in c for c in classes):
                break  # next section
            if sib.name != "table":
                continue
            for tr in sib.find_all("tr")[1:]:
                cells = tr.find_all(["td", "th"])
                if not cells:
                    continue
                name = cells[0].get_text(" ", strip=True)
                if len(name) > 5:
                    out.setdefault(_normalize(name), current)
    return out


# ── НЦТ XLSX: programs + offerings + thresholds (joined to egov) ──────────────
def parse_xlsx(path, egov_index, ovpo_acc):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    programs, offerings, thresholds = {}, {}, []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row or len(row) < 10:
            continue
        ovpo, uni_name, form = row[2], row[3], row[4]
        area_name, gop_code, gop_name, score = row[6], row[7], row[8], row[9]
        if not (ovpo and uni_name and gop_code):
            continue
        gop = _norm_code(str(gop_code))
        if not re.fullmatch(r"BM?\d{3}", gop):
            continue
        if form and "полная" not in str(form).lower():
            continue

        # resolve university: match XLSX name -> egov id, else ovpo fallback
        uid = _best_match(_normalize(str(uni_name)), egov_index, threshold=0.6)
        if uid is None:
            uid = f"ovpo-{str(ovpo).strip()}"
            oblast = str(row[1]).strip() if row[1] else ""
            if oblast.upper().startswith("ГОРОД "):
                oblast = oblast[6:].strip()
            ovpo_acc.setdefault(uid, {
                "id": uid, "name_ru": str(uni_name).strip(),
                "city": oblast.title() or None, "source_url": XLSX_URL,
            })

        programs.setdefault(gop, {
            "code": gop,
            "name_ru": str(gop_name).strip() if gop_name else gop,
            "field": _field_for(str(area_name) if area_name else None),
            "source_url": XLSX_URL,
        })
        offerings.setdefault((uid, gop), {
            "university_id": uid, "program_code": gop, "source_url": XLSX_URL,
        })
        if isinstance(score, (int, float)):
            thresholds.append({
                "program_code": gop, "university_id": uid, "year": GRANT_YEAR,
                "quota_type": "general", "metric": "competition_min",
                "min_score": int(score), "is_verified": True,
                "source_url": XLSX_URL,
                "note": "Минимальный балл для допуска к конкурсу на грант, "
                        "очная полная форма (НЦТ, 2024)",
            })
    return programs, offerings, thresholds


def parse_pdf_subjects(path):
    reader = pypdf.PdfReader(path)
    text = re.sub(r"\s+", " ", " ".join(p.extract_text() or "" for p in reader.pages))
    matches = list(re.compile(r"[ВB]М?\d{3}").finditer(text))
    out = {}
    for i, m in enumerate(matches):
        code = _norm_code(m.group())
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        chunk = text[m.end():end]
        found = sorted(
            (sm.start(), subj)
            for subj in SUBJECTS
            for sm in re.finditer(re.escape(subj), chunk)
        )
        if len(found) >= 2:
            out[code] = (found[-2][1], found[-1][1])
    return out


def main():
    ap = argparse.ArgumentParser()
    here = os.path.dirname(os.path.abspath(__file__))
    ap.add_argument("--out", default=os.path.join(
        here, "..", "..", "assets", "data", "universities.json"))
    ap.add_argument("--egov")
    ap.add_argument("--wiki")
    ap.add_argument("--xlsx")
    ap.add_argument("--pdf")
    args = ap.parse_args()

    egov_path = args.egov or _download(EGOV_URL, ".html")
    wiki_path = args.wiki or _download(WIKI_URL, ".html")
    xlsx_path = args.xlsx or _download(XLSX_URL, ".xlsx")
    pdf_path = args.pdf or _download(PDF_URL, ".pdf")

    universities = parse_egov(egov_path)
    wiki_types = parse_wiki_types(wiki_path)
    egov_index = {u["id"]: _normalize(u["name_ru"]) for u in universities}

    # merge type onto egov unis (confident match only)
    typed = 0
    for u in universities:
        toks = egov_index[u["id"]]
        best, bj = None, 0.0
        for nt, val in wiki_types.items():
            j = _jaccard(toks, nt)
            if j > bj:
                bj, best = j, val
        if best and bj >= 0.6:
            u["type"] = best
            typed += 1

    ovpo_acc = {}
    programs, offerings, thresholds = parse_xlsx(xlsx_path, egov_index, ovpo_acc)
    subjects = parse_pdf_subjects(pdf_path)

    enriched = 0
    for code, prog in programs.items():
        if code in subjects:
            prog["ent_profile_subject_1"], prog["ent_profile_subject_2"] = subjects[code]
            prog["subjects_source_url"] = PDF_URL
            enriched += 1

    universities.extend(ovpo_acc.values())
    matched = len({o["university_id"] for o in offerings.values()
                   if not o["university_id"].startswith("ovpo-")})

    out = {
        "meta": {
            "schema_version": 1,
            "generated_at": _dt.date.today().isoformat(),
            "generated_by": "tools/scrape_universities/scrape.py",
            "note": "competition_min = минимум для допуска к конкурсу на грант "
                    "(НЦТ 2024), НЕ проходной балл. Вузы — реестр egov.kz, "
                    "тип — Wikipedia (по совпадению имени), баллы — НЦТ XLSX, "
                    "профильные предметы — НЦТ PDF.",
            "sources": {
                "registry_egov": EGOV_URL, "types_wikipedia": WIKI_URL,
                "grant_scores_xlsx": XLSX_URL, "gop_subjects_pdf": PDF_URL,
            },
        },
        "universities": sorted(universities, key=lambda u: u["id"]),
        "education_programs": sorted(programs.values(), key=lambda p: p["code"]),
        "university_programs": sorted(
            offerings.values(), key=lambda o: (o["university_id"], o["program_code"])),
        "grant_thresholds": thresholds,
    }
    out_path = os.path.abspath(args.out)
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=2)

    print(
        f"WROTE {out_path}\n"
        f"  universities={len(out['universities'])} "
        f"(egov+{len(ovpo_acc)} ovpo-fallback, typed={typed})\n"
        f"  programs={len(out['education_programs'])} (subjects on {enriched})\n"
        f"  offerings={len(out['university_programs'])} "
        f"(grant unis matched to egov={matched})\n"
        f"  thresholds={len(out['grant_thresholds'])}",
        file=sys.stderr,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
